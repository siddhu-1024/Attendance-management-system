import io
import csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch


# =============================================================================
# CSV EXPORTERS
# =============================================================================

def generate_session_csv(session_data, records_list, college_name="Smart Classroom College of Engineering"):
    """
    Generate CSV for a single attendance session.
    """
    output = io.StringIO()
    writer = csv.writer(output)

    writer.writerow([college_name])
    writer.writerow(["ATTENDANCE SESSION REPORT"])
    writer.writerow([])
    writer.writerow(["Date:", session_data.get("date")])
    writer.writerow(["Subject:", session_data.get("subject_name")])
    writer.writerow(["Section:", session_data.get("section")])
    writer.writerow(["Period:", f"Period {session_data.get('period')}"])
    writer.writerow(["Faculty:", session_data.get("faculty_name")])
    writer.writerow(["Topic:", session_data.get("topic") or "N/A"])
    writer.writerow(["Total Students:", session_data.get("total_students")])
    writer.writerow(["Present Count:", session_data.get("present_count")])
    writer.writerow(["Absent Count:", session_data.get("absent_count")])
    writer.writerow(["Attendance Percentage:", f"{session_data.get('percentage')}%"])
    writer.writerow([])
    writer.writerow(["S.No", "Roll Number", "Student Name", "Status", "Remarks"])

    for idx, rec in enumerate(records_list, 1):
        writer.writerow([
            idx,
            rec.get("roll_no"),
            rec.get("student_name"),
            rec.get("status"),
            rec.get("remarks") or ""
        ])

    writer.writerow([])
    writer.writerow(["Report Generated On:", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])

    output.seek(0)
    # Return as bytes with UTF-8 BOM so Excel opens CSV without encoding issues
    return output.getvalue().encode('utf-8-sig')


def generate_cumulative_csv(students_summary, filters_info=None, college_name="Smart Classroom College of Engineering"):
    """
    Generate CSV for cumulative student attendance summary.
    """
    output = io.StringIO()
    writer = csv.writer(output)

    writer.writerow([college_name])
    writer.writerow(["CUMULATIVE STUDENT ATTENDANCE REPORT"])
    if filters_info:
        writer.writerow(["Filters:", filters_info])
    writer.writerow([])
    writer.writerow(["S.No", "Roll Number", "Student Name", "Total Classes", "Present", "Absent", "Attendance %", "Status"])

    for idx, s in enumerate(students_summary, 1):
        status_flag = "Low Attendance" if s.get("is_low") else "Normal"
        writer.writerow([
            idx,
            s.get("roll_no"),
            s.get("name"),
            s.get("total"),
            s.get("present"),
            s.get("absent"),
            f"{s.get('percentage')}%",
            status_flag
        ])

    writer.writerow([])
    writer.writerow(["Report Generated On:", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])

    output.seek(0)
    return output.getvalue().encode('utf-8-sig')


# =============================================================================
# EXCEL (.XLSX) EXPORTERS
# =============================================================================

def generate_session_excel(session_data, records_list, college_name="Smart Classroom College of Engineering"):
    """
    Generate a beautifully styled Excel workbook for an attendance session.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Session"

    # Styling definitions
    font_title = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    font_sub = Font(name="Calibri", size=11, bold=True, color="475569")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)

    fill_header = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    fill_meta_lbl = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_present = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_absent = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    font_present = Font(name="Calibri", size=10, bold=True, color="166534")
    font_absent = Font(name="Calibri", size=10, bold=True, color="991B1B")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Title Banner
    ws.merge_cells("A1:E1")
    ws["A1"] = college_name
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:E2")
    ws["A2"] = "SESSION ATTENDANCE RECORD"
    ws["A2"].font = font_sub
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 18

    # Metadata Grid
    meta_rows = [
        ("Date:", session_data.get("date"), "Subject:", session_data.get("subject_name")),
        ("Section:", session_data.get("section"), "Period:", f"Period {session_data.get('period')}"),
        ("Faculty:", session_data.get("faculty_name"), "Topic:", session_data.get("topic") or "N/A"),
        ("Total Students:", session_data.get("total_students"), "Attendance %:", f"{session_data.get('percentage')}%"),
        ("Present Count:", session_data.get("present_count"), "Absent Count:", session_data.get("absent_count")),
    ]

    curr_row = 4
    for r in meta_rows:
        ws.cell(row=curr_row, column=1, value=r[0]).font = font_bold
        ws.cell(row=curr_row, column=1).fill = fill_meta_lbl
        ws.cell(row=curr_row, column=2, value=r[1]).font = font_regular

        ws.cell(row=curr_row, column=4, value=r[2]).font = font_bold
        ws.cell(row=curr_row, column=4).fill = fill_meta_lbl
        ws.cell(row=curr_row, column=5, value=r[3]).font = font_regular
        curr_row += 1

    curr_row += 1

    # Student Table Headers
    headers = ["S.No", "Roll Number", "Student Name", "Attendance Status", "Remarks"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=curr_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center" if col_idx in [1, 4] else "left", vertical="center")
        cell.border = thin_border

    ws.row_dimensions[curr_row].height = 22
    curr_row += 1

    # Student Rows
    for idx, rec in enumerate(records_list, 1):
        c1 = ws.cell(row=curr_row, column=1, value=idx)
        c2 = ws.cell(row=curr_row, column=2, value=rec.get("roll_no"))
        c3 = ws.cell(row=curr_row, column=3, value=rec.get("student_name"))
        c4 = ws.cell(row=curr_row, column=4, value=rec.get("status"))
        c5 = ws.cell(row=curr_row, column=5, value=rec.get("remarks") or "")

        c1.alignment = Alignment(horizontal="center")
        c2.alignment = Alignment(horizontal="center")
        c4.alignment = Alignment(horizontal="center")

        is_present = rec.get("status") == "Present"
        c4.fill = fill_present if is_present else fill_absent
        c4.font = font_present if is_present else font_absent

        for cell in [c1, c2, c3, c5]:
            cell.font = font_regular
            cell.border = thin_border
        c4.border = thin_border

        curr_row += 1

    # Adjust Column Widths
    col_widths = {1: 8, 2: 18, 3: 38, 4: 20, 5: 25}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[openpyxl_get_col_letter(col_idx)].width = width

    # Write to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def generate_cumulative_excel(students_summary, filters_info=None, threshold=75.0, college_name="Smart Classroom College of Engineering"):
    """
    Generate styled Excel for cumulative student attendance.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Cumulative Attendance"

    font_title = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    font_sub = Font(name="Calibri", size=11, bold=True, color="475569")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_regular = Font(name="Calibri", size=10)
    font_bold = Font(name="Calibri", size=10, bold=True)

    fill_header = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    fill_low = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    font_low = Font(name="Calibri", size=10, bold=True, color="991B1B")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.merge_cells("A1:H1")
    ws["A1"] = college_name
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:H2")
    ws["A2"] = "CUMULATIVE ATTENDANCE SUMMARY"
    ws["A2"].font = font_sub
    ws["A2"].alignment = Alignment(horizontal="center")

    curr_row = 4
    if filters_info:
        ws.cell(row=curr_row, column=1, value=f"Filters: {filters_info}").font = font_bold
        curr_row += 1

    headers = ["S.No", "Roll Number", "Student Name", "Total Classes", "Present", "Absent", "Attendance %", "Standing"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=curr_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 4, 5, 6, 7, 8] else "left", vertical="center")
        cell.border = thin_border

    curr_row += 1

    for idx, s in enumerate(students_summary, 1):
        pct = float(s.get("percentage", 0))
        is_low = pct < threshold

        c1 = ws.cell(row=curr_row, column=1, value=idx)
        c2 = ws.cell(row=curr_row, column=2, value=s.get("roll_no"))
        c3 = ws.cell(row=curr_row, column=3, value=s.get("name"))
        c4 = ws.cell(row=curr_row, column=4, value=s.get("total"))
        c5 = ws.cell(row=curr_row, column=5, value=s.get("present"))
        c6 = ws.cell(row=curr_row, column=6, value=s.get("absent"))
        c7 = ws.cell(row=curr_row, column=7, value=f"{pct:.1f}%")
        c8 = ws.cell(row=curr_row, column=8, value="⚠️ Low Attendance" if is_low else "Normal")

        c1.alignment = Alignment(horizontal="center")
        c2.alignment = Alignment(horizontal="center")
        c4.alignment = Alignment(horizontal="center")
        c5.alignment = Alignment(horizontal="center")
        c6.alignment = Alignment(horizontal="center")
        c7.alignment = Alignment(horizontal="center")
        c8.alignment = Alignment(horizontal="center")

        for cell in [c1, c2, c3, c4, c5, c6, c7, c8]:
            cell.font = font_low if is_low else font_regular
            cell.border = thin_border
            if is_low and cell in [c7, c8]:
                cell.fill = fill_low

        curr_row += 1

    col_widths = {1: 8, 2: 18, 3: 38, 4: 15, 5: 12, 6: 12, 7: 16, 8: 20}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[openpyxl_get_col_letter(col_idx)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def openpyxl_get_col_letter(col_idx):
    letters = 'ABCDEFGHIJKLMNOPQRSTUVWXYZ'
    result = ''
    while col_idx > 0:
        col_idx, remainder = divmod(col_idx - 1, 26)
        result = letters[remainder] + result
    return result


# =============================================================================
# PDF EXPORTERS (ReportLab)
# =============================================================================

def generate_session_pdf(session_data, records_list, college_name="Smart Classroom College of Engineering"):
    """
    Generate a professional PDF report for an attendance session.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        alignment=1, # Center
        textColor=colors.HexColor('#1E3A8A')
    )

    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        alignment=1,
        textColor=colors.HexColor('#475569')
    )

    meta_label_style = ParagraphStyle(
        'MetaLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#1E293B')
    )

    meta_val_style = ParagraphStyle(
        'MetaVal',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#334155')
    )

    th_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11,
        alignment=1,
        textColor=colors.white
    )

    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#1E293B')
    )

    cell_center = ParagraphStyle(
        'TableCellCenter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        alignment=1,
        textColor=colors.HexColor('#1E293B')
    )

    present_style = ParagraphStyle(
        'PresentBadge',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        alignment=1,
        textColor=colors.HexColor('#166534')
    )

    absent_style = ParagraphStyle(
        'AbsentBadge',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        alignment=1,
        textColor=colors.HexColor('#991B1B')
    )

    elements = []

    # Header
    elements.append(Paragraph(college_name.upper(), title_style))
    elements.append(Spacer(1, 4))
    elements.append(Paragraph("OFFICIAL ATTENDANCE SESSION REPORT", subtitle_style))
    elements.append(Spacer(1, 12))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#1E40AF'), spaceAfter=10))

    # Session Metadata Box
    date_val = session_data.get("date", "")
    subject_val = session_data.get("subject_name", "")
    section_val = session_data.get("section", "")
    period_val = f"Period {session_data.get('period', '1')}"
    faculty_val = session_data.get("faculty_name", "Faculty")
    topic_val = session_data.get("topic") or "Class Lecture"
    
    total_val = str(session_data.get("total_students", 0))
    present_val = str(session_data.get("present_count", 0))
    absent_val = str(session_data.get("absent_count", 0))
    pct_val = f"{session_data.get('percentage', 0.0)}%"

    meta_table_data = [
        [
            Paragraph("Date:", meta_label_style), Paragraph(date_val, meta_val_style),
            Paragraph("Subject:", meta_label_style), Paragraph(subject_val, meta_val_style)
        ],
        [
            Paragraph("Section:", meta_label_style), Paragraph(section_val, meta_val_style),
            Paragraph("Period:", meta_label_style), Paragraph(period_val, meta_val_style)
        ],
        [
            Paragraph("Faculty:", meta_label_style), Paragraph(faculty_val, meta_val_style),
            Paragraph("Topic:", meta_label_style), Paragraph(topic_val, meta_val_style)
        ],
        [
            Paragraph("Total Students:", meta_label_style), Paragraph(total_val, meta_val_style),
            Paragraph("Attendance %:", meta_label_style), Paragraph(pct_val, meta_val_style)
        ],
        [
            Paragraph("Present Count:", meta_label_style), Paragraph(present_val, meta_val_style),
            Paragraph("Absent Count:", meta_label_style), Paragraph(absent_val, meta_val_style)
        ]
    ]

    meta_table = Table(meta_table_data, colWidths=[100, 160, 100, 160])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))

    elements.append(meta_table)
    elements.append(Spacer(1, 14))

    # Student Roster Table
    table_data = [[
        Paragraph("S.No", th_style),
        Paragraph("Roll Number", th_style),
        Paragraph("Student Name", th_style),
        Paragraph("Status", th_style),
        Paragraph("Remarks", th_style)
    ]]

    for idx, rec in enumerate(records_list, 1):
        is_p = rec.get("status") == "Present"
        status_p = Paragraph(rec.get("status"), present_style if is_p else absent_style)
        
        table_data.append([
            Paragraph(str(idx), cell_center),
            Paragraph(rec.get("roll_no"), cell_center),
            Paragraph(rec.get("student_name"), cell_style),
            status_p,
            Paragraph(rec.get("remarks") or "-", cell_style)
        ])

    roster_table = Table(table_data, colWidths=[35, 95, 230, 75, 85], repeatRows=1)
    
    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E40AF')),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 3.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
    ]

    # Alternate row colors
    for i in range(1, len(table_data)):
        if i % 2 == 0:
            t_style.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor('#F8FAFC')))

    roster_table.setStyle(TableStyle(t_style))
    elements.append(roster_table)

    elements.append(Spacer(1, 25))

    # Signatures
    sig_data = [
        [
            Paragraph("________________________<br/>Faculty Signature", meta_label_style),
            Paragraph("________________________<br/>HOD / Dean Signature", meta_label_style)
        ]
    ]
    sig_table = Table(sig_data, colWidths=[260, 260])
    sig_table.setStyle(TableStyle([
        ('ALIGN', (0, 0), (0, 0), 'LEFT'),
        ('ALIGN', (1, 0), (1, 0), 'RIGHT'),
    ]))
    elements.append(KeepTogether(sig_table))

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()


def generate_cumulative_pdf(students_summary, filters_info=None, threshold=75.0, college_name="Smart Classroom College of Engineering"):
    """
    Generate PDF for cumulative student attendance.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        alignment=1,
        textColor=colors.HexColor('#1E3A8A')
    )

    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10.5,
        leading=13,
        alignment=1,
        textColor=colors.HexColor('#475569')
    )

    th_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=10,
        alignment=1,
        textColor=colors.white
    )

    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1E293B')
    )

    cell_center = ParagraphStyle(
        'TableCellCenter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        alignment=1,
        textColor=colors.HexColor('#1E293B')
    )

    low_badge = ParagraphStyle(
        'LowBadge',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        alignment=1,
        textColor=colors.HexColor('#991B1B')
    )

    good_badge = ParagraphStyle(
        'GoodBadge',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        alignment=1,
        textColor=colors.HexColor('#166534')
    )

    elements = []

    elements.append(Paragraph(college_name.upper(), title_style))
    elements.append(Spacer(1, 4))
    elements.append(Paragraph("CUMULATIVE ATTENDANCE & DEFAULTER SUMMARY", subtitle_style))
    if filters_info:
        elements.append(Spacer(1, 2))
        elements.append(Paragraph(f"Criteria: {filters_info} | Low Attendance Threshold: {threshold}%", cell_center))
    elements.append(Spacer(1, 10))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#1E40AF'), spaceAfter=10))

    table_data = [[
        Paragraph("S.No", th_style),
        Paragraph("Roll No", th_style),
        Paragraph("Student Name", th_style),
        Paragraph("Classes", th_style),
        Paragraph("Present", th_style),
        Paragraph("Absent", th_style),
        Paragraph("%", th_style),
        Paragraph("Standing", th_style)
    ]]

    for idx, s in enumerate(students_summary, 1):
        pct = float(s.get("percentage", 0))
        is_low = pct < threshold

        pct_p = Paragraph(f"{pct:.1f}%", low_badge if is_low else good_badge)
        status_p = Paragraph("Low Attendance" if is_low else "Regular", low_badge if is_low else good_badge)

        table_data.append([
            Paragraph(str(idx), cell_center),
            Paragraph(s.get("roll_no"), cell_center),
            Paragraph(s.get("name"), cell_style),
            Paragraph(str(s.get("total")), cell_center),
            Paragraph(str(s.get("present")), cell_center),
            Paragraph(str(s.get("absent")), cell_center),
            pct_p,
            status_p
        ])

    roster_table = Table(table_data, colWidths=[30, 85, 185, 45, 45, 45, 45, 40], repeatRows=1)
    
    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E40AF')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]

    for i in range(1, len(table_data)):
        row_student = students_summary[i - 1]
        pct = float(row_student.get("percentage", 0))
        if pct < threshold:
            t_style.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor('#FEE2E2')))
        elif i % 2 == 0:
            t_style.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor('#F8FAFC')))

    roster_table.setStyle(TableStyle(t_style))
    elements.append(roster_table)

    elements.append(Spacer(1, 20))

    sig_data = [
        [
            Paragraph("________________________<br/>Class Incharge", cell_style),
            Paragraph("________________________<br/>Head of Department", cell_style)
        ]
    ]
    sig_table = Table(sig_data, colWidths=[260, 260])
    sig_table.setStyle(TableStyle([
        ('ALIGN', (0, 0), (0, 0), 'LEFT'),
        ('ALIGN', (1, 0), (1, 0), 'RIGHT'),
    ]))
    elements.append(KeepTogether(sig_table))

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()
